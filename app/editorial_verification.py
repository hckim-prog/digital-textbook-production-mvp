"""Real GUI publication with fixture responses; never contacts an API."""
import json
from pathlib import Path
import shutil
import time
from unittest.mock import patch
from uuid import uuid4

from docx import Document
from PIL import Image
from PySide6.QtCore import QSettings
from PySide6.QtWidgets import QScrollArea

from app.controllers.production import Production
from app.controllers.publishing import publish
from core.ai import service


def run(app, root):
    import app.gui.window as gui
    base = root / 'reports/ed-v'
    sandbox = base / uuid4().hex[:6]
    shutil.copytree(root / 'config', sandbox / 'config')
    source = sandbox / 'source.docx'
    text = '함수는 이름을 통해 호출하며 입력을 받아 결과를 반환합니다.'
    doc = Document()
    for value in ['Chapter 1', 'Chapter 1 함수', '1.1 함수 호출', text]:
        doc.add_paragraph(value)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'int main() {\n\treturn 0;\n}'
    picture = sandbox / 'image.png'
    Image.new('RGB',(60,40),'steelblue').save(picture)
    doc.add_picture(str(picture))
    doc.add_paragraph('연습문제'); doc.add_paragraph('(추후 제공)'); doc.save(source)
    original = source.read_bytes()
    job = Production(sandbox, source); master = job.analyze()
    # Reproduce an already accepted outline containing a title-only chapter,
    # as in the saved depth analysis that reached learning generation.
    from core.manuscript.structure import propose_rules
    job.structure().save(propose_rules(master), approved=True)
    report = {'passed':False, 'api_mode':'fixture_no_paid_calls', 'sandbox':str(sandbox)}
    calls, errors = [], []
    def response(root, model, effort, prompt):
        p = json.loads(prompt.split('자료:\n',1)[1]); calls.append(p)
        if 'questions' in p:
            d = {'accepted':True,'reason':'고정 시험의 원고 근거·정답 검수'}
        elif 'sources' in p:
            b = next(row for row in p['sources'] if row['text']==text)
            prompts = ['함수는 무엇을 통해 호출하는가?', '함수는 무엇을 입력으로 받는가?', '함수가 반환하는 것은 무엇인가?',
                       '함수의 입력과 반환 결과를 구분해 설명하라.', '함수 호출에서 이름의 역할을 설명하라.']
            d = {'questions':[{'question':q,'answer':'함수는 이름으로 호출하고 입력을 받아 결과를 반환한다.',
                              'explanation':text,'evidence':[{'block_id':b['id'],'quote':text}]}
                             for q in prompts[:2 if p['type']=='section' else 5]]}
        elif 'units' in p:
            d = {'issues':[]}
        elif 'technical_change' in p:
            d = {'accept':True,'meaning_preserved':True,'evidence_supported':False,'reason':'의미 유지'}
        else:
            d = {'text':p['candidate'],'technical_change':False,'evidence':[],'reason':'유지'}
        return d, {'model':model}
    with patch.object(service,'_response',response), patch.object(service,'record',lambda *a:None), \
         patch.object(gui,'QSettings',lambda *_:QSettings(str(sandbox/'prefs.ini'),QSettings.IniFormat)), \
         patch.object(gui.QMessageBox,'question',lambda *a:gui.QMessageBox.Yes), \
         patch.object(gui.QMessageBox,'warning',lambda *a:errors.append(str(a[-1]))):
        win = gui.MainWindow(sandbox)
        try:
            win.depth_existing.setChecked(True)
            win.source_edit.setText(str(source)); win.prepare_timer.stop(); win._analysis_done(job, master)
            assert win.learning_ai.isChecked() and win.editorial_ai.isChecked()
            win.show(); win.build()
            deadline = time.monotonic() + 40
            while (win.busy or any(t.isRunning() for t in win.threads)) and time.monotonic() < deadline:
                app.processEvents(); time.sleep(.01)
            assert not win.busy and not any(t.isRunning() for t in win.threads), 'GUI timeout'
            assert not errors, errors
            built = job.last_result()
            assert built and built['qa']['passed'] and built['learning']['question_count']==7
            assert [u['title'] for u in built['learning']['skipped_units']]==['Chapter 1']
            assert not any(p.get('id')=='outline-0001' for p in calls if 'sources' in p)
            assert built['editorial']['ai_audited'] and source.read_bytes()==original
            assert win.editorial_button.isEnabled() and '추가 학습 문제: 7개' in win.result_info.text()
            scroll = win.findChild(QScrollArea); scroll.ensureWidgetVisible(win.learning_ai)
            app.processEvents(); win.grab().save(str(base/'gui.png'))
            before = job.last_result()['folder']
            try:
                publish(job,['web'])  # Original placeholder must still be blocked without enrichment.
            except ValueError as exc:
                assert '완결성 검사 미통과' in str(exc)
            else:
                raise AssertionError('placeholder accepted')
            assert job.last_result()['folder']==before
            report.update(passed=True, gui_options=True, all_formats=True, source_preserved=True, title_only_skipped=True,
                          placeholder_blocked=True, last_success_preserved=True,
                          requests=len(calls), output=built['folder'])
        except Exception as exc:
            report['error']=str(exc)
        finally:
            if any(t.isRunning() for t in win.threads):
                win.cancel_review.set()
                while any(t.isRunning() for t in win.threads):
                    app.processEvents(); time.sleep(.01)
            win.close()
    (base/'acceptance.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0 if report['passed'] else 1
