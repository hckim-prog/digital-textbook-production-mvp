"""Deterministic hierarchy GUI/export acceptance check; never calls an API."""
import json
from pathlib import Path
from uuid import uuid4

from docx import Document
from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QMessageBox


def run(app, root):
    import app.gui.structure_dialog as gui
    from app.controllers.production import Production
    from core.manuscript.structure import propose_rules
    audit = root / "reports/structure-verification"
    sandbox = audit / ("run-" + uuid4().hex[:8])
    sandbox.mkdir(parents=True)
    source = sandbox / "구조 검증.docx"
    doc = Document()
    doc.add_paragraph("이 장에서는 프로그램의 구조를 배웁니다.")
    doc.add_heading("1.1. 프로그램 이해", level=2)
    doc.add_heading("1.1.1. 코드 살펴보기", level=3)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '#include <iostream>\n\nint main() {\n\treturn 0;\n}'
    doc.add_heading("1.2. 실행 결과", level=2)
    doc.add_paragraph("화면에서 실행 결과를 확인합니다.")
    doc.save(source)
    job = Production(sandbox / "project", source)
    job.analyze()
    dialog = gui.StructureDialog(job, "검증용 고정 응답", "none")
    report = {"api_mode": "fixture_no_paid_calls", "passed": False, "sandbox": str(sandbox)}
    state = {"phase": 0, "ticks": 0, "failure_seen": False}
    gui.propose_ai = lambda root, master, *_: propose_rules(master)
    QMessageBox.question = lambda *_: QMessageBox.Yes
    QMessageBox.warning = lambda *_: state.update(failure_seen=True)
    dialog.show()
    def finish():
        (audit / "acceptance.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        timer.stop()
        app.exit(0 if report["passed"] else 1)
    def tick():
        state["ticks"] += 1
        if state["ticks"] > 100:
            report["error"] = "verification timeout"
            finish()
            return
        try:
            if state["phase"] == 0:
                state["phase"] = 1
                dialog.ai()
            elif state["phase"] == 1 and dialog.thread is None:
                assert len(dialog.nodes) == 4
                dialog.tree.setCurrentItem(dialog.tree.topLevelItem(0).child(0).child(0))
                dialog.grab().save(str(audit / "structure-dialog.png"))
                def failure(*_):
                    raise ValueError("검증용 API 실패")
                gui.propose_ai = failure
                state["phase"] = 2
                dialog.ai()
            elif state["phase"] == 2 and dialog.thread is None:
                assert state["failure_seen"] and len(dialog.nodes) == 4
                dialog.approve()
                result = job.build(["web", "pdf", "epub"])
                report.update(passed=result["qa"]["passed"], qa=result["qa"],
                              ai_success_and_failure_ui=True, result_folder=result["folder"])
                finish()
        except Exception as exc:
            report["error"] = str(exc)
            finish()
    timer = QTimer()
    timer.timeout.connect(tick)
    timer.start(150)
    return app.exec()
