"""Opt-in frozen UI integration test with deterministic AI responses, no API cost."""
import json
import shutil
import sys
from pathlib import Path
from uuid import uuid4

from docx import Document
from PySide6.QtCore import QSettings, QTimer
from PySide6.QtWidgets import QMessageBox, QScrollArea

from core.ai import workflow as engine


def run(app, root):
    import app.gui.window as gui
    audit = root / 'reports/review-upgrade'
    sandbox = audit / ('frozen-' + uuid4().hex[:8])
    shutil.copytree(root / 'config', sandbox / 'config')
    source = sandbox / '순차 검토 시험.docx'
    doc = Document()
    doc.add_paragraph('처음 문장입니다.')
    doc.add_paragraph('다른 문장입니다.')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '#include <iostream>\n\nint main() {\n\treturn 0;\n}'
    doc.save(source)
    gui.QSettings = lambda *_: QSettings(str(sandbox / 'prefs.ini'), QSettings.IniFormat)
    win = gui.MainWindow(sandbox)
    win.depth_existing.setChecked(True)
    win.setWindowTitle('검토 흐름 검증 · 테스트 응답 사용')
    win.quick_mode.setChecked(False)
    win.quick_ai.setChecked(False)
    win.detail_toggle.setChecked(True)
    win.show()
    state = {'stage': 0}
    report = {'frozen': bool(getattr(sys, 'frozen', False)), 'executable': sys.executable,
              'api_mode': 'deterministic_fixture_no_paid_calls', 'inputs': [], 'sandbox': str(sandbox)}
    targets = ['교정한 문장입니다.', '교정하고 기술을 확인한 문장입니다.', '교정하고 기술과 최종 표현을 확인한 문장입니다.']

    def fake_ai(root, master, model, reasoning, limit, progress, role, cancelled=None):
        report['inputs'].append({'role': role, 'text': master.blocks[0].text})
        items = [{'id': 'new', 'role': role, 'block_id': master.blocks[0].id,
                  'source_text': master.blocks[0].text, 'suggested_text': targets[engine.ROLES.index(role)],
                  'level': 'A', 'reason': '단계 연결 검증용 응답', 'status': 'pending', 'model': model}]
        calls = [{'success': True, 'block_id': b.id, 'input_tokens': 0, 'output_tokens': 0,
                  'elapsed_seconds': 0, 'estimated_cost_usd': 0} for b in master.blocks]
        return items, calls
    engine.proofread = fake_ai
    QMessageBox.question = lambda *args: QMessageBox.Yes
    def fail(message):
        report['error'] = str(message)
        win.busy = False
        state['stage'] = -1
    win._failed = fail
    QMessageBox.warning = lambda *args: fail(args[-1])

    def finish(passed):
        report['passed'] = passed
        (audit / 'exe-flow.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        win.verification_timer.stop()
        win.close()
        app.exit(0 if passed else 1)

    def tick():
        try:
            if win.busy or win.prepare_timer.isActive() or any(t.isRunning() for t in win.threads):
                return
            stage = state['stage']
            if stage == -1:
                finish(False)
            elif stage == 0:
                assert win.scope.currentData() == 'all'
                assert '이번 대상 2개' in win.scope_info.text()
                win.grab().save(str(audit / 'light-ui-top.png'))
                win.proofread_button.click()
                state['stage'] = 1
            elif stage in (1, 2, 3):
                assert win.table.rowCount() == 1
                win.table.selectRow(0)
                win.review_buttons[0].click()
                if stage < 3:
                    win.active_role.setCurrentIndex(stage)
                    assert win.table.rowCount() == 0
                    win.proofread_button.click()
                    state['stage'] += 1
                else:
                    assert [i['text'] for i in report['inputs']] == ['처음 문장입니다.', targets[0], targets[1]]
                    job = win.production()
                    assert [i['status'] for i in job.suggestions()] == ['approved'] * 3
                    assert job.workflow().output_master().blocks[0].text == targets[-1]
                    win.build_button.click()
                    state['stage'] = 4
            elif stage == 4:
                result = win.production().last_result()
                assert result['qa']['passed']
                report['result'] = result
                report['decisions'] = win.production().suggestions()
                win.review_filter.setCurrentIndex(1)
                assert win.table.rowCount() == 3
                scroll = win.findChild(QScrollArea)
                scroll.verticalScrollBar().setValue(scroll.verticalScrollBar().maximum())
                state['stage'] = 5
            elif stage == 5:
                win.grab().save(str(audit / 'exe-flow.png'))
                win.comparison_preview = win.suggestion_dialog(0)
                win.comparison_preview.show()
                state['stage'] = 6
            elif stage == 6:
                dialog = win.comparison_preview
                assert dialog.comparison_editors[0].toPlainText() == '처음 문장입니다.'
                assert dialog.comparison_editors[1].toPlainText() == targets[0]
                dialog.grab().save(str(audit / 'light-ui-comparison.png'))
                dialog.close()
                from PySide6.QtGui import QPalette
                assert app.palette().color(QPalette.Window).lightness() > 230
                report['light_theme'] = True
                report['comparison_text_preserved'] = True
                finish(True)
        except Exception as exc:
            fail(type(exc).__name__ + ': ' + str(exc))
    win.source_edit.setText(str(source))
    win.verification_timer = QTimer(win)
    win.verification_timer.timeout.connect(tick)
    win.verification_timer.start(100)
    app.verification_window = win
    return app.exec()
