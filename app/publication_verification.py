"""Frozen acceptance check using only local documents and fixture AI."""
import json
import shutil
from pathlib import Path
from uuid import uuid4
from docx import Document
from app.controllers.production import Production
from app.controllers import publishing
from core.manuscript.structure import node


def run(app, root):
    sandbox=root/'reports/publication-verification'/('run-'+uuid4().hex[:8])
    shutil.copytree(root/'config',sandbox/'config')
    result={'passed':False,'api_mode':'fixture_no_paid_calls'}
    try:
        source=sandbox/'structured.docx';doc=Document()
        for t in ['Chapter 1 시작','1.1 첫 번째 주제','첫 번째 본문입니다.','1.2 두 번째 주제','두 번째 본문입니다.']:doc.add_paragraph(t)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'int main() {\n\treturn 0;\n}'
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'cout << getLarger(10.0, 20.5) << endl;'
        doc.save(source)
        job=Production(sandbox,source);job.analyze()
        for theme in ['auto', 'standard', 'lab', 'reading']:
            built=publishing.publish(job,['web','pdf','epub'],theme=theme)
            assert built['qa']['passed']
            assert built['design']['id'] == ('lab' if theme == 'auto' else theme)
        result['design_themes']=['auto','standard','lab','reading']
        result['automatic_outline_output']=True
        from core.manuscript import depth
        def fake_depth(root, model, reasoning, prompt):
            rows = json.loads(prompt.split('Section 자료:\n', 1)[1])
            anchor = rows[1]['id']
            additions = [{'level':3,'start_block_id':anchor,'title':'시험 하위 주제','evidence':'고정 응답 검증'}]
            if '최대 뎁스는 4.' in prompt:
                additions.append({'level':4,'start_block_id':anchor,'title':'시험 세부 주제','evidence':'고정 응답 검증'})
            return {'nodes': additions}, {'model': model}
        depth._response = fake_depth
        depth.record = lambda *args: None
        for maximum in (3,4):
            expanded=publishing.publish(job,['web','pdf','epub'],depth_mode=True,max_depth=maximum,structure_model='fixture')
            assert expanded['qa']['passed']
            completed=json.loads((job.work/'ai-depth-final.json').read_text(encoding='utf-8'))
            assert max(n['level'] for n in completed['nodes']) == maximum
        result['ai_depth_direct_outputs'] = [3,4]
        plain=sandbox/'plain.docx';doc=Document();doc.add_paragraph('구조 없는 본문입니다.');doc.save(plain)
        raw=Production(sandbox,plain);raw.analyze()
        try:publishing.publish(raw,['web'])
        except ValueError:result['unstructured_blocked']=True
        else:raise AssertionError('Unstructured source was accepted')
        publishing.propose_ai=lambda root,m,*a,**k:[node(m.blocks[0],1,'장 제목'),node(m.blocks[0],2,'절 제목')]
        ai=publishing.publish(raw,['web','pdf','epub'],allow_restructure=True,structure_model='fixture')
        assert ai['qa']['passed']
        result.update(passed=True,ai_opt_in_output=True)
    except Exception as exc:
        result['error']=str(exc)
    target=root/'reports/publication-verification/acceptance.json'
    target.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0 if result['passed'] else 1
