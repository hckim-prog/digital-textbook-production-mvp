from pathlib import Path

import pytest
from docx import Document

from app.controllers.production import Production
from core.master import file_hash


def source_docx(tmp_path):
    folder = tmp_path / "원본"
    folder.mkdir()
    source = folder / "Chapter 1 원고.docx"
    doc = Document()
    doc.add_heading("Chapter 1", 0)
    doc.add_paragraph("이 문장을 교정합니다.")
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "#include <iostream>\n\nint main() {\n    return 0;\n}"
    doc.save(source)
    return source


@pytest.mark.parametrize("formats", [
    ["web"], ["pdf"], ["epub"], ["web", "pdf", "epub"],
])
def test_selected_output_folder_and_formats(tmp_path, formats):
    root = tmp_path / "project"
    root.mkdir()
    source = source_docx(tmp_path)
    before = file_hash(source)
    destination = tmp_path / "한글 결과물"
    destination.mkdir()
    job = Production(root, source)
    job.analyze()
    result = job.build(formats, destination)
    folder = Path(result["folder"])
    assert folder.parent == destination
    assert set(result["outputs"]) == set(formats)
    assert all(Path(path).is_file() for path in result["outputs"].values())
    assert (folder / "reports" / "code-fidelity.json").is_file()
    assert (folder / "reports" / "QA-report.html").is_file()
    assert (folder / "HTML" / "index.html").exists() == ("web" in formats)
    assert result["qa"]["passed"]
    assert file_hash(source) == before
    for kind in ("pdf", "epub"):
        if kind in formats:
            assert Path(result["outputs"][kind]).stem == source.stem


def test_default_folder_and_existing_job_are_not_overwritten(tmp_path):
    root = tmp_path / "project"
    root.mkdir()
    source = source_docx(tmp_path)
    job = Production(root, source)
    job.analyze()
    first = job.build(["web"])
    page = Path(first["outputs"]["web"])
    marker = page.read_bytes()
    second = job.build(["web"])
    assert Path(first["folder"]).parent == root / "output"
    assert first["folder"] != second["folder"]
    assert page.read_bytes() == marker


def test_missing_and_unwritable_destination(tmp_path, monkeypatch):
    root = tmp_path / "project"
    root.mkdir()
    source = source_docx(tmp_path)
    job = Production(root, source)
    job.analyze()
    with pytest.raises(RuntimeError, match="저장 폴더"):
        job.build(["web"], tmp_path / "없음")
    base = tmp_path / "읽기 전용"
    base.mkdir()
    import app.controllers.production as production
    def denied(*args, **kwargs):
        raise PermissionError("denied")
    monkeypatch.setattr(production.tempfile, "NamedTemporaryFile", denied)
    with pytest.raises(RuntimeError, match="파일을 저장"):
        job.build(["web"], base)


def test_gui_models_workflow_and_remembered_folder(tmp_path, monkeypatch):
    import os
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication
    import app.gui.window as window

    app = QApplication.instance() or QApplication([])
    import shutil
    actual_root = tmp_path / "gui-project"
    shutil.copytree(Path(__file__).resolve().parents[1] / "config", actual_root / "config")
    prefs_path = tmp_path / "prefs.ini"
    monkeypatch.setattr(window, "QSettings", lambda *_: QSettings(str(prefs_path), QSettings.IniFormat))
    win = window.MainWindow(actual_root)
    win.source_edit.setText("")
    assert not win.qa_button.isEnabled()
    assert not win.proofread_button.isEnabled()
    assert not win.build_button.isEnabled()
    assert all(box.count() == 6 for box in win.model_boxes.values())
    assert win.model_boxes["proofreading"].currentData() == "gpt-5.6-terra"
    assert win.model_boxes["technical_review"].currentData() == "gpt-5.6-sol"
    assert win.model_boxes["final_review"].currentData() == "gpt-5.6-sol"
    astra = win.model_boxes["proofreading"].findData("gpt-6-astra")
    win.model_boxes["proofreading"].setCurrentIndex(astra)
    assert [win.reasoning.itemText(i) for i in range(win.reasoning.count())] == ["low", "medium", "high", "xhigh", "max"]
    win.reasoning_role.setCurrentIndex(1)
    assert "none" in [win.reasoning.itemText(i) for i in range(win.reasoning.count())]
    win.reasoning_role.setCurrentIndex(0)
    source = source_docx(tmp_path)
    win.source_edit.setText(str(source))
    assert win.prepare_timer.isActive()
    assert not win.proofread_button.isEnabled()
    win._analysis_done(Production(actual_root, source), Production(actual_root, source).analyze())
    assert win.proofread_button.isEnabled()
    assert win.build_button.isEnabled()
    assert not any(button.isEnabled() for button in win.review_buttons)
    destination = tmp_path / "선택한 결과"
    destination.mkdir()
    monkeypatch.setattr(window.QFileDialog, "getExistingDirectory", lambda *args: str(destination))
    win.choose_output()
    assert win.output_edit.text() == str(destination)
    second = window.MainWindow(actual_root)
    assert second.output_edit.text() == str(destination)
    win.close()
    second.close()
