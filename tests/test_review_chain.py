import copy
import json
from pathlib import Path

from docx import Document
import pytest

from app.controllers.production import Production
from core.ai import workflow as engine


@pytest.fixture
def job(tmp_path):
    source = tmp_path / '원고.docx'
    doc = Document()
    doc.add_paragraph('처음 문장입니다.')
    doc.add_paragraph('다른 문장입니다.')
    doc.save(source)
    result = Production(tmp_path / 'project', source)
    result.analyze()
    return result


def fake_engine(monkeypatch, targets, received=None, fail_at=None):
    def run(root, master, model, reasoning, limit, progress, role, cancelled=None):
        items, calls = [], []
        if received is not None:
            received.append((role, [b.text for b in master.blocks]))
        for n, block in enumerate(master.blocks):
            failed = n == fail_at
            calls.append({'success': not failed, 'block_id': block.id, 'input_tokens': 1,
                          'output_tokens': 1, 'estimated_cost_usd': .001})
            if failed:
                break
            target = targets.get((role, block.id), targets.get(role))
            if target:
                items.append({'id': 'temporary', 'block_id': block.id, 'source_text': block.text,
                              'suggested_text': target, 'level': 'A', 'reason': '검증',
                              'model': model, 'role': role, 'status': 'pending'})
        return items, calls
    monkeypatch.setattr(engine, 'proofread', run)


def test_stage_inputs_include_only_earlier_approvals_and_final_output(job, monkeypatch):
    received = []
    fake_engine(monkeypatch, {'proofreading': '교정한 문장입니다.',
        'technical_review': '교정하고 기술을 확인한 문장입니다.',
        'final_review': '교정하고 기술과 최종 표현을 확인한 문장입니다.'}, received)
    for role in engine.ROLES:
        result = job.proofread('m', 'none', 1, role=role)
        assert result['success_count'] == 1
        item = next(i for i in job.suggestions() if i['role'] == role)
        job.review([item['id']], 'approved')
    assert received == [('proofreading', ['처음 문장입니다.']),
        ('technical_review', ['교정한 문장입니다.']),
        ('final_review', ['교정하고 기술을 확인한 문장입니다.'])]
    assert [i['status'] for i in job.suggestions()] == ['approved'] * 3
    assert job.workflow().output_master().blocks[0].text == '교정하고 기술과 최종 표현을 확인한 문장입니다.'
    restored = Production(job.root, job.source)
    result = restored.build(['web'])
    assert result['qa']['passed']
    assert '교정하고 기술과 최종 표현을 확인한 문장입니다.' in Path(result['outputs']['web']).read_text(encoding='utf-8')


def test_earlier_changes_invalidate_downstream_without_auto_hold(job, monkeypatch):
    fake_engine(monkeypatch, {'proofreading': '교정한 문장입니다.', 'technical_review': '기술 검토한 문장입니다.'})
    job.proofread('m','none',1)
    first = job.suggestions()[0]
    job.review([first['id']], 'approved')
    job.proofread('m','none',1,role='technical_review')
    second = job.suggestions()[1]
    job.review([second['id']], 'approved')
    job.review([first['id']], 'rejected')
    assert [i['status'] for i in job.suggestions()] == ['rejected','outdated']
    assert job.workflow().output_master().blocks[0].text == '처음 문장입니다.'
    with pytest.raises(ValueError, match='기준 문장'):
        job.review([second['id']], 'approved')
    # Reverting an upstream decision must not leave all old proposals stale and skipped.
    job.review([first['id']], 'approved')
    assert job.review_plan('m','none',1,role='technical_review')['target_count'] == 1


def test_same_stage_alternatives_are_explicitly_replaced(job, monkeypatch):
    fake_engine(monkeypatch, {'proofreading': '첫 교정안입니다.'})
    job.proofread('m','none',1)
    items = job.suggestions()
    second = {**copy.deepcopy(items[0]), 'id':'alternative', 'suggested_text':'대안 문장입니다.'}
    job.save_suggestions(items + [second])
    with pytest.raises(ValueError, match='동시에 승인'):
        job.review([items[0]['id'], 'alternative'], 'approved')
    assert all(i['status']=='pending' for i in job.suggestions())
    job.review([items[0]['id']], 'approved')
    job.review(['alternative'], 'approved')
    assert [i['status'] for i in job.suggestions()] == ['superseded', 'approved']
    job.review(['alternative'], 'hold')
    assert job.suggestions()[1]['status']=='hold'
    assert job.suggestions()[0]['history'][-1]['actor']=='automatic'


def test_legacy_decisions_are_preserved_and_approval_becomes_baseline(job, monkeypatch):
    block = job.master().blocks[0]
    legacy = [{'id':'old','block_id':block.id,'source_text':block.text,
               'suggested_text':'이미 승인한 문장입니다.','status':'approved','level':'A','reason':'기존 기록'}]
    job.save_suggestions(legacy)
    original_bytes = (job.work/'suggestions.json').read_bytes()
    received = []
    fake_engine(monkeypatch, {}, received)
    job.proofread('m','none',1,role='technical_review')
    assert received[0][1] == ['이미 승인한 문장입니다.']
    assert (job.work/'suggestions.json').read_bytes() == original_bytes
    assert job.workflow().legacy() == legacy
    assert job.workflow().output_master().blocks[0].text == '이미 승인한 문장입니다.'


def test_sampling_continues_and_failure_does_not_skip_remaining(job, monkeypatch):
    fake_engine(monkeypatch, {})
    first=job.proofread('m','none',1)
    assert first['success_count']==1
    plan=job.review_plan('m','none',1)
    assert plan['already_reviewed']==1 and plan['blocks'][0].text=='다른 문장입니다.'
    fake_engine(monkeypatch, {}, fail_at=0)
    failed=job.proofread('m','none',None)
    assert failed['failure_count']==1 and failed['success_count']==0
    assert job.review_plan('m','none')['target_count']==1
    fake_engine(monkeypatch, {})
    job.proofread('m','none')
    assert job.review_plan('m','none')['target_count']==0
    assert job.review_plan('m','none',repeat=True)['target_count']==2


def test_optional_stages_can_be_skipped(job, monkeypatch):
    fake_engine(monkeypatch, {'final_review':'최종 문장입니다.'})
    job.proofread('m','none',1,role='final_review')
    job.review([job.suggestions()[0]['id']], 'approved')
    assert job.workflow().output_master().blocks[0].text == '최종 문장입니다.'


def test_cancellation_stops_new_requests_and_can_resume(job, monkeypatch):
    from threading import Event
    from core.ai import service
    stop = Event()
    seen = []
    def response(root, model, reasoning, prompt):
        seen.append(prompt)
        stop.set()
        return {'suggestions': []}, {'model': model, 'input_tokens': 1, 'output_tokens': 1, 'elapsed_seconds': .1}
    monkeypatch.setattr(service, '_response', response)
    monkeypatch.setattr(service, 'record', lambda *a: None)
    monkeypatch.setattr(service, 'cost', lambda *a: .001)
    result = job.proofread('m','none', cancelled=stop.is_set)
    assert len(seen)==1 and result['cancelled']
    assert result['success_count']==1 and result['unprocessed_count']==1
    assert job.review_plan('m','none')['target_count']==1


def test_gui_stage_filter_scope_and_partial_failure(job, tmp_path, monkeypatch):
    import os
    import shutil
    import time
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication, QMessageBox
    import app.gui.window as gui
    app = QApplication.instance() or QApplication([])
    shutil.copytree(Path(__file__).resolve().parents[1] / 'config', job.root / 'config')
    monkeypatch.setattr(gui, 'QSettings', lambda *_: QSettings(str(tmp_path/'ui.ini'), QSettings.IniFormat))
    monkeypatch.setattr(QMessageBox, 'question', lambda *a: QMessageBox.Yes)
    win = gui.MainWindow(job.root)
    win.source_edit.setText(str(job.source))
    def wait():
        end = time.monotonic() + 8
        while win.busy or win.prepare_timer.isActive() or any(t.isRunning() for t in win.threads):
            app.processEvents()
            assert time.monotonic()<end
            time.sleep(.01)
        app.processEvents()
    wait()
    assert win.scope.currentData()=='all' and '이번 대상 2개' in win.scope_info.text()
    fake_engine(monkeypatch, {'proofreading':'수정 문장입니다.'})
    win.scope.setCurrentIndex(1)
    win.limit.setValue(1)
    win.proofread_button.click(); wait()
    assert win.table.rowCount()==1 and win.table.item(0,4).text()=='A 단순 교정'
    win.table.selectRow(0); win.review_buttons[0].click()
    win.active_role.setCurrentIndex(1)
    assert win.table.rowCount()==0
    fake_engine(monkeypatch, {'technical_review':'기술 수정 문장입니다.'}, fail_at=1)
    win.scope.setCurrentIndex(0)
    win.proofread_button.click(); wait()
    assert '일부 실패' in win.status.text() and win.progress.value()==50
    assert '실패 1' in win.status.text()
    assert win.table.rowCount()==1 and win.table.item(0,1).text()=='수정 문장입니다.'
    win.review_filter.setCurrentIndex(1)
    assert win.table.rowCount()==2
    win.close()
