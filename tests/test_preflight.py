import copy
from pathlib import Path
import pytest
from docx import Document
from core.master import Master, Block
from core.manuscript.preflight import inspect_structure
from core.manuscript.structure import node
from app.controllers.production import Production
from app.controllers import publishing


def make_job(tmp_path, structured=True):
    p=tmp_path/'source.docx';doc=Document()
    if structured:
        doc.add_paragraph('Chapter 1 시작')
        doc.add_paragraph('1.1 첫 번째 주제')
    doc.add_paragraph('본문을 읽어 봅시다.')
    if structured:
        doc.add_paragraph('1.2 두 번째 주제')
        doc.add_paragraph('다른 내용입니다.')
    doc.save(p)
    job=Production(tmp_path/'project',p);job.analyze()
    return job


def test_unstyled_numbered_titles_accepted(tmp_path):
    job=make_job(tmp_path)
    r=publishing.preflight(job)
    assert r['passed'] and len(r['nodes'])==3


def test_no_headings_blocks_before_any_api(tmp_path,monkeypatch):
    job=make_job(tmp_path,False)
    monkeypatch.setattr(job,'proofread',lambda *a,**k:pytest.fail('Paid proofreading reached'))
    monkeypatch.setattr(publishing,'propose_ai',lambda *a,**k:pytest.fail('Paid structure reached'))
    with pytest.raises(ValueError,match='구조 보완 필요'):
        publishing.publish(job,['web'],run_ai=True)
    assert not (job.work/'last-result.json').exists()


def test_number_conflict_and_skipped_level():
    m=Master('t','s','h',[Block('a','paragraph','Chapter 1 시작'),Block('b','paragraph','2.1 잘못된 번호'),Block('c','paragraph','본문')])
    assert any(i['code']=='number_conflict' for i in inspect_structure(m)['issues'])
    m.blocks[1].text='1.1.1 누락된 상위 제목'
    assert any(i['code']=='omitted_heading' for i in inspect_structure(m)['issues'])


def test_list_numbers_do_not_become_titles():
    m=Master('t','s','h',[Block('a','procedure-step','1. 설치한다'),Block('b','procedure-step','2. 실행한다')])
    assert not inspect_structure(m)['passed']


def test_long_unsegmented_and_prose_style_flagged():
    blocks=[Block('h','paragraph','Chapter 1 시작')]+[Block(str(i),'paragraph','본문입니다.') for i in range(41)]
    assert any(i['code']=='long_unsegmented_range' for i in inspect_structure(Master('t','s','h',blocks))['issues'])
    blocks[0]=Block('h','heading','이것은 일반적인 본문입니다.',style='Heading 1')
    assert any(i['code']=='prose_heading' for i in inspect_structure(Master('t','s','h',blocks))['issues'])


def test_auto_export_keeps_draft_and_source_untouched(tmp_path):
    job=make_job(tmp_path);original=job.source.read_bytes()
    job.structure().save([node(job.master().blocks[0],1)],approved=False)
    draft=(job.work/'structure.json').read_bytes()
    r=publishing.publish(job,['web','pdf','epub'])
    assert r['qa']['passed']
    assert job.source.read_bytes()==original and (job.work/'structure.json').read_bytes()==draft
    assert (Path(r['folder'])/'reports/publication-preflight.json').is_file()
    assert len(Master.load(Path(r['folder'])/'reports/approved-master.json').outline)==3


def test_ai_opt_in_is_validated_then_used_without_approval(tmp_path,monkeypatch):
    job=make_job(tmp_path,False);master=job.master();calls=[]
    def fake(root,m,model,reasoning,mode):
        calls.append((model,reasoning,mode))
        return [node(m.blocks[0],1,'새 장'),node(m.blocks[0],2,'새 절')]
    monkeypatch.setattr(publishing,'propose_ai',fake)
    r=publishing.publish(job,['web'],allow_restructure=True,structure_model='chosen',structure_reasoning='low')
    assert r['qa']['passed'] and calls==[('chosen','low','create')]
    publishing.publish(job,['web'],allow_restructure=True,structure_model='chosen',structure_reasoning='low')
    assert len(calls)==1
    assert not (job.work/'structure.json').exists()
    assert job.master()==master


def test_invalid_ai_structure_stops_before_proofreading(tmp_path,monkeypatch):
    job=make_job(tmp_path,False)
    monkeypatch.setattr(publishing,'propose_ai',lambda root,m,*a,**k:[node(m.blocks[0],1,'임시 장')])
    monkeypatch.setattr(job,'proofread',lambda *a,**k:pytest.fail('Proofreading reached'))
    with pytest.raises(ValueError,match='AI 구조 검사 미통과'):
        publishing.publish(job,['web'],run_ai=True,allow_restructure=True,structure_model='chosen')


def test_changed_source_blocks_even_with_ai_opt_in(tmp_path,monkeypatch):
    job=make_job(tmp_path);doc=Document(job.source);doc.add_paragraph('추가');doc.save(job.source)
    monkeypatch.setattr(publishing,'propose_ai',lambda *a,**k:pytest.fail('AI reached'))
    with pytest.raises(ValueError,match='원고 보완 필요'):
        publishing.publish(job,['web'],allow_restructure=True)


def test_output_failure_not_reported_as_success(tmp_path,monkeypatch):
    job=make_job(tmp_path)
    monkeypatch.setattr(job,'build',lambda *a,**k:{'qa':{'passed':False},'report':'qa.html'})
    with pytest.raises(RuntimeError,match='출력 보존 검사 미통과'):
        publishing.publish(job,['web'])


def test_duplicate_chapter_and_empty_body(tmp_path):
    m=Master('t','s','h',[Block('a','paragraph','Chapter 1 시작'),Block('b','paragraph','내용'),Block('c','paragraph','Chapter 1 반복')])
    assert any(i['code']=='chapter_order' for i in inspect_structure(m)['issues'])
    doc=Document();doc.add_paragraph('Chapter 1 시작');p=tmp_path/'empty.docx';doc.save(p)
    job=Production(tmp_path/'project',p);job.analyze()
    assert any(i['code']=='no_body' for i in publishing.preflight(job)['hard_issues'])
