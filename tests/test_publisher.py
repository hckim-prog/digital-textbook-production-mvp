import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import pymupdf

from app.controllers.production import Production
from core.qa.code_fidelity import check_output_codes
from core.reporting.comparison import refresh, review


def fixture_job(tmp_path):
    doc = Document()
    doc.add_heading('💡 원고와 숫자 001010', 1)
    doc.add_paragraph('원문 문장입니다.')
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = '#include <iostream>\n\nint main() {\n\tint n = 1;\n    // < > &\n\n    return 0;\n}'
    image = tmp_path / 'sample.png'
    Image.new('RGB', (40, 30), 'red').save(image)
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = '이미지 설명'
    cell.paragraphs[0].add_run().add_picture(str(image))
    paragraph = doc.add_paragraph('수식 ')
    container = OxmlElement('m:oMathPara')
    for number in ('1', '2'):
        math = OxmlElement('m:oMath')
        run = OxmlElement('m:r')
        text = OxmlElement('m:t')
        text.text = number
        run.append(text); math.append(run); container.append(math)
    paragraph._p.append(container)
    source = tmp_path / 'source.docx'
    doc.save(source)
    root = tmp_path / 'project'
    root.mkdir()
    job = Production(root, source)
    job.analyze()
    return job


def test_rich_content_and_real_code_outputs(tmp_path):
    job = fixture_job(tmp_path)
    result = job.build(['web', 'pdf', 'epub'])
    assert result['qa']['passed'], result['qa']
    assert result['qa']['image_count'] == 1
    assert result['qa']['warnings']  # Formula layout needs human review.
    paths = {key: Path(path) for key, path in result['outputs'].items()}
    assert check_output_codes(job.source, job.master(), paths)['passed']
    with pymupdf.open(paths['pdf']) as doc:
        assert '💡' in ''.join(page.get_text() for page in doc)
        evidence = json.loads(doc.embfile_get('code-layout.json'))
        line = next(line for line in evidence if 'int n' in line['text'])
        page = doc[line['page']]
        page.add_redact_annot(pymupdf.Rect(line['x'], line['baseline'] - line['size'], line['x'] + line['width'], line['baseline'] + 1))
        page.apply_redactions()
        doc.saveIncr()
    # Embedded raw source alone must not let a visibly damaged PDF pass.
    assert not check_output_codes(job.source, job.master(), {'pdf': paths['pdf']})['passed']
    with ZipFile(paths['epub']) as book:
        files = {name: book.read(name) for name in book.namelist()}
    files['EPUB/chapter.xhtml'] = files['EPUB/chapter.xhtml'].replace(b'\tint n', b' int n')
    with ZipFile(paths['epub'], 'w', ZIP_DEFLATED) as book:
        for name, data in files.items():
            book.writestr(name, data)
    assert not check_output_codes(job.source, job.master(), {'epub': paths['epub']})['passed']


def test_approval_persists_and_only_approved_text_is_exported(tmp_path):
    job = fixture_job(tmp_path)
    block = next(b for b in job.master().blocks if b.text == '원문 문장입니다.')
    items = [{'id': f's{i}', 'block_id': block.id, 'source_text': block.text,
              'suggested_text': target, 'status': 'pending'} for i, target in enumerate(['승인 문장입니다.', '거절 문장입니다.', '보류 문장입니다.'])]
    job.workflow().start()
    for item in items:
        item['role'] = 'proofreading'
    job.save_suggestions(items)
    for key, state in [('s0','approved'),('s1','rejected'),('s2','hold')]:
        job.review([key], state)
    restored = Production(job.root, job.source)
    assert [i['status'] for i in restored.suggestions()] == ['approved','rejected','hold']
    result = restored.build(['web'])
    text = Path(result['outputs']['web']).read_text(encoding='utf-8')
    assert '승인 문장입니다.' in text and '거절 문장입니다.' not in text and '보류 문장입니다.' not in text
    assert result['qa']['passed']
    assert restored.last_result()['folder'] == result['folder']


def test_comparison_rate_uses_reviewed_decisions(tmp_path):
    row = dict(model_id='model', suggestions=3, elapsed_seconds=1, input_tokens=2, output_tokens=3, estimated_cost_usd=.01)
    (tmp_path / 'model-comparison-rows.json').write_text(json.dumps([row]), encoding='utf-8')
    items = [{'id': str(i), 'status': 'pending'} for i in range(3)]
    (tmp_path / 'model-comparison-details.json').write_text(json.dumps([{'model_id':'model','suggestions':items}]), encoding='utf-8')
    refresh(tmp_path)
    def current():
        return json.loads((tmp_path/'model-comparison-rows.json').read_text(encoding='utf-8'))[0]
    assert current()['approval_rate'] == '미검토'
    review(tmp_path, [('model','0')], 'approved')
    review(tmp_path, [('model','1')], 'rejected')
    review(tmp_path, [('model','2')], 'hold')
    assert current()['approval_rate'] == 50 and current()['hold'] == 1


def test_code_paragraph_with_image_does_not_hang(tmp_path):
    image = tmp_path / 'sample.png'
    Image.new('RGB',(4,4),'blue').save(image)
    doc = Document()
    p = doc.add_paragraph('int main() {\n    return 0;\n}')
    p.add_run().add_picture(str(image))
    file = tmp_path/'code.docx'; doc.save(file)
    from core.manuscript.docx_reader import read_docx
    master = read_docx(file,tmp_path/'work')
    assert master.blocks[0].text == 'int main() {\n    return 0;\n}'
    assert len(master.blocks[0].assets) == 1
