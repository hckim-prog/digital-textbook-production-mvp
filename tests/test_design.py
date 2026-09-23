import copy
import json
from dataclasses import asdict
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import html

from app.controllers.production import Production
from app.controllers.publishing import publish
from core.export.preview import sample_master, create_preview
from core.export.themes import resolve, THEMES
from core.export.pdf_layout import CodeBlock
from core.export.outputs import _pdf_fonts
from core.master import Master, Block, file_hash


@pytest.mark.parametrize('theme', list(THEMES))
def test_theme_exports_preserve_source_and_code(tmp_path, theme):
    source = tmp_path / '원고.docx'
    doc = Document()
    doc.add_paragraph('Chapter 1 시작')
    doc.add_paragraph('1.1 입력과 출력')
    doc.add_paragraph('값 1234와 문자열을 확인합니다.')
    code = 'int main() {\n\n\tint x = 1234;  \n\treturn x;\n}'
    doc.add_table(rows=1, cols=1).cell(0, 0).text = code
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip([c for row in table.rows for c in row.cells], ['항목','값','정수','1234']):
        cell.text = text
    doc.add_paragraph('1.2 다음 주제')
    doc.add_paragraph('두 번째 절의 본문입니다.')
    doc.save(source)
    digest = file_hash(source)
    job = Production(tmp_path / 'project', source)
    master = job.analyze()
    before = copy.deepcopy(asdict(master))
    built = publish(job, ['web', 'pdf', 'epub'], theme=theme)
    assert built['qa']['passed'], built['qa']['checks']
    assert built['design']['id'] == theme
    assert json.loads((job.reports / 'design.json').read_text(encoding='utf-8'))['requested'] == theme
    assert file_hash(source) == digest
    assert asdict(job.master()) == before
    tree = html.parse(built['outputs']['web'])
    assert tree.xpath('//pre/code')[0].text == code
    assert tree.xpath('//meta[@name="book-theme"]/@content') == [theme]
    assert tree.xpath('//aside//details/summary')[0].text == '목차'
    assert not tree.xpath('//script[@src] | //link[@rel="stylesheet"]')
    with ZipFile(built['outputs']['epub']) as archive:
        chapter = html.fromstring(archive.read('EPUB/chapter.xhtml'))
        assert not chapter.xpath('//script | //button')
        assert chapter.xpath('//pre/code')[0].text == code


def test_theme_auto_rules_and_preview(tmp_path):
    master = sample_master()
    before = asdict(master)
    assert resolve(master)['id'] == 'lab'
    assert resolve(Master('t','s','h',[Block('b','paragraph','본문')]))['id'] == 'standard'
    assert resolve(Master('t','s','h',[Block('b','paragraph','본문' * 5000)]))['id'] == 'reading'
    assert asdict(master) == before
    assert create_preview(tmp_path).is_file()
    with pytest.raises(ValueError):
        resolve(master, 'invalid')


def test_pdf_long_code_never_silently_shrinks_to_unreadable_size():
    _pdf_fonts()
    with pytest.raises(ValueError, match='너무 길어'):
        CodeBlock('long', 'x' * 1000, []).wrap(473, 700)


def test_preview_without_manuscript_and_saved_selection(tmp_path, monkeypatch):
    import shutil
    from pathlib import Path
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication
    import app.gui.window as gui
    shutil.copytree(Path(__file__).resolve().parents[1] / 'config', tmp_path / 'config')
    monkeypatch.setattr(gui, 'QSettings', lambda *_: QSettings(str(tmp_path / 'prefs.ini'), QSettings.IniFormat))
    opened = []
    monkeypatch.setattr(gui.QDesktopServices, 'openUrl', lambda url: opened.append(url.toLocalFile()) or True)
    def no_warning(*args):
        pytest.fail('Preview should work without selecting a manuscript')
    monkeypatch.setattr(gui.QMessageBox, 'warning', no_warning)
    app = QApplication.instance() or QApplication([])
    win = gui.MainWindow(tmp_path)
    win.preview_design()
    assert Path(opened[-1]).is_file()
    win.design_theme.setCurrentIndex(win.design_theme.findData('reading'))
    win.preview_design()
    assert html.parse(opened[-1]).xpath('//meta[@name="book-theme"]/@content') == ['reading']
    win.close()
    again = gui.MainWindow(tmp_path)
    assert again.design_theme.currentData() == 'reading'
    again.close()
