"""Exercise the simplified UI's saved choices and paid-action boundary."""
import json
import os
import shutil
from pathlib import Path

from docx import Document

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')


def make_window(tmp_path, monkeypatch):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication
    import app.gui.window as gui
    app = QApplication.instance() or QApplication([])
    root = tmp_path / 'project'
    shutil.copytree(Path(__file__).resolve().parents[1] / 'config', root / 'config')
    prefs = QSettings(str(tmp_path / 'prefs.ini'), QSettings.IniFormat)
    prefs.setValue('proofreading', 'gpt-6-sol')
    prefs.setValue('technical_review', 'gpt-6-astra')
    prefs.setValue('depth_mode', 'existing')
    monkeypatch.setattr(gui, 'QSettings', lambda *_: prefs)
    return app, gui.MainWindow(root), prefs


def test_main_page_keeps_saved_models_and_depth_choices(tmp_path, monkeypatch):
    app, win, prefs = make_window(tmp_path, monkeypatch)
    try:
        win.show(); app.processEvents()
        assert win.pages.currentIndex() == 0
        assert win.quick_ai.isVisible() and win.learning_ai.isVisible()
        assert win.depth_option.isVisible() and not win.depth_option.isChecked()
        assert not win.model_boxes['proofreading'].isVisible()
        win.depth_option.setChecked(True)
        assert win.depth_ai.isChecked() and prefs.value('depth_mode') == 'ai'
        win.pages.setCurrentIndex(2); app.processEvents()
        assert win.model_boxes['proofreading'].isVisible()
        assert win.model_boxes['proofreading'].currentData() == 'gpt-6-sol'
        assert win.model_boxes['technical_review'].currentData() == 'gpt-6-astra'
        win.depth_existing.setChecked(True)
        assert not win.depth_option.isChecked()
        win.pages.setCurrentIndex(0)
        assert 'GPT-6 Astra' in win.settings_summary.text()
        win.result_info.setText('\n'.join(['제작 완료', '저장 폴더: 시험', '디자인: 표준', '최종 검사: 확인 필요', '추가 학습 문제: 7개', '자동 수정: 3건']))
        win._fit_result_info()
        for _ in range(4):
            app.processEvents()
        assert win.result_info.height() >= win.result_info.heightForWidth(win.result_info.width())
        assert win.qa_button.y() >= win.result_info.y() + win.result_info.height()
    finally:
        win.close()


def test_structure_only_ai_requires_cost_confirmation(tmp_path, monkeypatch):
    import app.gui.window as gui
    app, win, _ = make_window(tmp_path, monkeypatch)
    from app.controllers.production import Production
    source = tmp_path / 'source.docx'
    doc = Document(); doc.add_paragraph('Chapter 1 함수'); doc.add_paragraph('1.1 함수 호출')
    doc.add_paragraph('함수는 이름으로 호출합니다.'); doc.save(source)
    job = Production(win.root, source)
    win.source_edit.setText(str(source)); win.prepare_timer.stop()
    win._analysis_done(job, job.analyze())
    win.quick_ai.setChecked(False)
    prompts, starts = [], []
    monkeypatch.setattr(gui.QMessageBox, 'question', lambda *args: (prompts.append(args[-1]), gui.QMessageBox.No)[1])
    monkeypatch.setattr(win, 'run_task', lambda *args: starts.append(args))
    try:
        win.depth_option.setChecked(True); win.build()
        assert len(prompts) == 1 and '소제목 자동 보완' in prompts[-1] and not starts
        win.depth_option.setChecked(False); win.allow_restructure.setChecked(True); win.build()
        assert len(prompts) == 2 and '새 목차 구성' in prompts[-1] and not starts
        win.allow_restructure.setChecked(False); win.build()
        assert len(prompts) == 2 and len(starts) == 1
        assert not job.last_result() and source.exists()
    finally:
        win.close()


def test_comparison_displays_model_with_no_suggestions(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication
    from app.gui.comparison import ComparisonDialog
    from core.reporting import comparison
    from core.master import Block, Master
    app = QApplication.instance() or QApplication([])
    source = '할수 있다.'
    def response(root, sample, model, effort, limit):
        items = [{'id': 's1', 'level': 'A', 'source_text': source, 'suggested_text': '할 수 있다.', 'reason': '띄어쓰기'}] if model == 'sol' else []
        return items, [{'response_model': model, 'success': True, 'estimated_cost_usd': 0}]
    monkeypatch.setattr(comparison, 'proofread', response)
    page = comparison.compare(tmp_path, Master('시험', 'sample.docx', 'hash', [Block('b1', 'paragraph', source)]), ['sol', 'luna'], 'low', 5)
    dialog = ComparisonDialog(page.parent)
    try:
        assert dialog.table.rowCount() == 1 and dialog.model_results.rowCount() == 2
        assert dialog.model_results.item(1, 0).text() == 'luna'
        assert dialog.model_results.item(1, 1).text() == '검토 완료'
        assert dialog.model_results.item(1, 3).text() == '0건'
        assert dialog.model_results.item(1, 4).text() == '제안 없음'
        dialog.table.selectRow(0); dialog.decide('approved')
        assert dialog.model_results.item(0, 4).text() == '100.0%'
        rows = json.loads((page.parent / 'model-comparison-rows.json').read_text(encoding='utf-8'))
        assert rows[1]['suggestions'] == 0
    finally:
        dialog.close()
