import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import html
from PIL import Image

from app.controllers.production import Production
from app.controllers.publishing import preflight, publish, diagnostic_summary
from core.manuscript.content import table_parts
from core.master import file_hash
from core.qa.checks import check
from core.qa.code_fidelity import check_output_codes


def nested_job(tmp_path):
    doc = Document()
    doc.add_heading('Chapter 1 표', 1)
    doc.add_heading('1.1 내부 표', 2)
    doc.add_paragraph('표 안의 작은 표와 앞뒤 설명을 보존합니다.')
    outer = doc.add_table(rows=1, cols=1).cell(0, 0)
    outer.text = '앞 설명 001'
    table = outer.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '이진수 001010001'
    code = '#include <iostream>\n\nint main() {\n\tint n = 1;\n    return 0;\n}'
    table.cell(0, 1).text = code
    p = table.cell(1, 0).paragraphs[0]
    p.add_run('수식 ')
    math = OxmlElement('m:oMath')
    run = OxmlElement('m:r'); text = OxmlElement('m:t'); text.text = 'x+1'
    run.append(text); math.append(run); p._p.append(math)
    image = tmp_path / 'sample.png'
    Image.new('RGB', (20, 15), 'blue').save(image)
    p.add_run().add_picture(str(image))
    table.cell(1, 1).text = '더 작은 표'
    deep = table.cell(1, 1).add_table(rows=1, cols=2)
    deep.cell(0, 0).text = 'A 01'
    deep.cell(0, 1).text = 'B 10'
    p = table.cell(1, 1).add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), doc.part.relate_to('https://example.com/reference', RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '참고 링크'
    run.append(text); link.append(run); p._p.append(link)
    outer.add_paragraph('뒤 설명 010')
    p = doc.add_paragraph('수학 첨자 ᵀ ᵢ ⱼ 그리고 ')
    p.add_run('\uf02d').font.name = 'Symbol'
    p.add_run(' 마지막\u200b')
    doc.add_paragraph('일반 문장의 끝 V\u200b')
    jpeg = tmp_path / 'sample.jpeg'
    picture = Image.new('RGB', (40, 30))
    picture.putdata([(x * 7 % 256, x * 13 % 256, x * 21 % 256) for x in range(1200)])
    picture.save(jpeg, quality=85)
    doc.add_paragraph().add_run().add_picture(str(jpeg))
    source = tmp_path / 'source.docx'; doc.save(source)
    root = tmp_path / 'project'; root.mkdir()
    job = Production(root, source); job.analyze()
    return job, code


def test_nested_tables_preserve_rich_content_grid_and_exact_code(tmp_path):
    job, code = nested_job(tmp_path)
    before = file_hash(job.source)
    assert not preflight(job)['hard_issues']
    master = job.master()
    block = next(b for b in master.blocks if b.kind == 'table')
    parts = list(table_parts(block.id, block.rows, block.cell_kinds, block.rich_cells))
    assert [list(map(len, rows)) for _, rows, _, _ in parts] == [[1], [2, 2], [2]]
    assert parts[1][1][0][1] == code
    from core.qa.editorial import text_units
    units = list(text_units(master))
    assert not any('#include' in unit['text'] or 'int main' in unit['text'] for unit in units)
    assert any('앞 설명' in unit['text'] for unit in units) and any('뒤 설명' in unit['text'] for unit in units)
    assert block.rows[0][0].index('앞 설명') < block.rows[0][0].index('001010001') < block.rows[0][0].index('뒤 설명')
    result = job.build(['web', 'pdf', 'epub'])
    assert result['qa']['passed'], result['qa']
    outputs = {kind: Path(path) for kind, path in result['outputs'].items()}
    assert check_output_codes(job.source, master, outputs)['passed']
    assert file_hash(job.source) == before
    assert result['qa']['image_count'] == 2 and result['qa']['hyperlink_count'] == 1
    original_html = outputs['web'].read_text(encoding='utf-8')
    outputs['web'].write_text(original_html.replace('\tint n', ' int n'), encoding='utf-8')
    assert not check_output_codes(job.source, master, {'web': outputs['web']})['passed']
    outputs['web'].write_text(original_html, encoding='utf-8')
    # A drawing-trace check must reject a truly missing zero-width character.
    import copy
    from core.export.outputs import pdf
    damaged = copy.deepcopy(master)
    for b in damaged.blocks:
        b.text = b.text.replace('\u200b', '')
        for item in b.inlines:
            item['text'] = item.get('text', '').replace('\u200b', '')
    broken_pdf = pdf(damaged, job.work / 'assets', tmp_path / 'damaged-pdf')
    assert not check(master, job.source, job.work / 'assets', {'pdf': broken_pdf})['passed']
    tree = html.parse(str(outputs['web']))
    # Joining cells into one paragraph retains text but must fail structure QA.
    nested = tree.xpath('//table[@data-table-id]')[1]
    nested.tag = 'div'
    tree.write(str(outputs['web']), encoding='utf-8', method='html')
    assert not check(master, job.source, job.work / 'assets', {'web': outputs['web']})['passed']


def test_missing_nested_content_blocks_before_ai_with_short_message(tmp_path, monkeypatch):
    import pytest
    job, _ = nested_job(tmp_path)
    master = job.master()
    block = next(b for b in master.blocks if b.kind == 'table')
    nested = next(item for item in block.rich_cells[0][0] if item['kind'] == 'table')
    block.rich_cells[0][0].remove(nested)
    block.rows[0][0] = '앞 설명 001\n뒤 설명 010'
    master.save(job.work / 'structured-master.json')
    monkeypatch.setattr(job, 'proofread', lambda *a, **k: pytest.fail('paid AI must not start'))
    with pytest.raises(ValueError) as exc:
        publish(job, ['web'], run_ai=True, model='test')
    assert len(str(exc.value)) < 800
    assert 'AI 문장 교정 미실행' in str(exc.value) and '진단 결과 보기' in str(exc.value)
    report = json.loads((job.work / 'preflight.json').read_text(encoding='utf-8'))
    assert report['hard_issues'] and not job.last_result()
    assert '원문 보존 확인 필요' in diagnostic_summary(report)


def test_old_review_baseline_restores_protected_table_without_changing_records(tmp_path):
    from core.master import Master
    from core.ai.workflow import write_json
    job, _ = nested_job(tmp_path)
    workflow = job.workflow(); workflow.start()
    baseline_path = workflow.folder / 'baseline.json'
    baseline = Master.load(baseline_path)
    block = next(b for b in baseline.blocks if b.kind == 'table')
    block.rich_cells[0][0] = [{'kind': 'text', 'text': '앞 설명 001\n뒤 설명 010'}]
    block.rows[0][0] = '앞 설명 001\n뒤 설명 010'
    block.assets = []
    # A skipped nested image also gave the following image an older filename.
    figure = next(b for b in baseline.blocks if b.kind == 'figure')
    figure.assets = ['image-0001.jpeg']
    for item in figure.inlines:
        if item['kind'] == 'image':
            item['asset'] = 'image-0001.jpeg'
    baseline.save(baseline_path)
    paragraph = next(b for b in baseline.blocks if b.text == '표 안의 작은 표와 앞뒤 설명을 보존합니다.')
    approved = paragraph.text.replace('보존', '유지')
    records = [{'id': f's{i}', 'block_id': paragraph.id, 'source_text': paragraph.text,
                'suggested_text': approved, 'role': 'proofreading', 'status': status,
                'history': [{'actor': 'user', 'to': status}]} for i, status in enumerate(['approved','rejected','hold'])]
    write_json(workflow.folder / 'suggestions.json', records)
    before = {p: file_hash(p) for p in workflow.folder.rglob('*') if p.is_file()}
    result = job.build(['web', 'pdf', 'epub'])
    assert result['qa']['passed'], result['qa']
    assert approved in Path(result['outputs']['web']).read_text(encoding='utf-8')
    assert all(file_hash(p) == digest for p, digest in before.items())
    assert workflow.items() == records
