import json
from pathlib import Path

import pytest
from docx import Document

from app.controllers.production import Production
from app.controllers.publishing import publish
from core.ai import automatic, service
from core.ai.workflow import write_json
from core.master import Master


MODELS = {role:(role,'none') for role in automatic.ROLES}


def job_for(tmp_path, texts=None):
    source = tmp_path / '자동출판.docx'
    doc = Document()
    for text in ['Chapter 1 시작','1.1 본문', *(texts or ['실제 호출할 함수를 결정합니다.'])]:
        doc.add_paragraph(text)
    doc.add_table(rows=1,cols=1).cell(0,0).text = 'int main() {\n\treturn 0;\n}'
    doc.save(source)
    job = Production(tmp_path/'project',source)
    job.analyze()
    return job


def fixtures(monkeypatch, response=None):
    calls=[]
    monkeypatch.setattr(service,'model_config',lambda *args:{'enabled':True,'reasoning_options':['none']})
    monkeypatch.setattr(service,'record',lambda *args:None)
    def fake(root, model, effort, prompt):
        payload=json.loads(prompt.split('자료:\n',1)[1])
        calls.append((model,payload))
        if response:
            data=response(model,payload)
        elif model=='final_review':
            data={'accept':True,'meaning_preserved':True,'evidence_supported':False,'reason':'의미 유지 확인'}
        else:
            data={'text':payload['candidate'].replace('실제 호출','실제로 호출'),
                  'technical_change':False,'evidence':[],'reason':'표현 교정'}
        return data,{'model':model}
    monkeypatch.setattr(service,'_response',fake)
    return calls


def test_automatic_three_outputs_cache_and_rollback(tmp_path,monkeypatch):
    job=job_for(tmp_path)
    original=job.source.read_bytes()
    calls=fixtures(monkeypatch)
    first=publish(job,['web','pdf','epub'],run_ai=True,automatic_models=MODELS)
    assert first['qa']['passed']
    assert first['quick']['applied_count']==1
    assert [m for m,_ in calls]==list(automatic.ROLES)
    assert all('int main' not in json.dumps(p,ensure_ascii=False) for _,p in calls)
    out=Master.load(Path(first['folder'])/'reports/approved-master.json')
    assert any(b.text=='실제로 호출할 함수를 결정합니다.' for b in out.blocks)
    publish(job,['web'],run_ai=True,automatic_models=MODELS)
    assert len(calls)==3
    write_json(job.work/'automatic-exclusions.json',[first['quick']['applied'][0]['block_id']])
    third=publish(job,['web'],run_ai=True,automatic_models=MODELS)
    assert third['quick']['applied_count']==0 and len(calls)==3
    reverted=Master.load(Path(third['folder'])/'reports/approved-master.json')
    assert any(b.text=='실제 호출할 함수를 결정합니다.' for b in reverted.blocks)
    assert job.source.read_bytes()==original and job.workflow().items()==[]


@pytest.mark.parametrize('status',['approved','hold','rejected'])
def test_manual_decisions_are_locked(tmp_path,monkeypatch,status):
    job=job_for(tmp_path)
    workflow=job.workflow();workflow.start()
    block=job.master().blocks[2]
    write_json(workflow.folder/'suggestions.json',[{'id':'manual','block_id':block.id,
        'source_text':block.text,'suggested_text':block.text+' 검토.', 'status':status,'role':'proofreading'}])
    before=(workflow.folder/'suggestions.json').read_bytes()
    calls=fixtures(monkeypatch)
    _,report=automatic.run(job,MODELS)
    assert not calls and not report['applied']
    assert (workflow.folder/'suggestions.json').read_bytes()==before


@pytest.mark.parametrize('failure',['number','meaning','ungrounded'])
def test_uncertain_changes_preserve_source_without_blocking_publication(tmp_path,monkeypatch,failure):
    job=job_for(tmp_path,['기준 값은 10이며 함수 호출이 가능합니다.'])
    def response(role,p):
        if role=='final_review':
            return {'accept':False,'meaning_preserved':False,'evidence_supported':False,'reason':'근거 부족'}
        target=p['original'].replace('10','20') if failure=='number' else p['original'].replace('가능','불가능')
        return {'text':target,'technical_change':failure=='ungrounded','evidence':[], 'reason':'시험'}
    fixtures(monkeypatch,response)
    result=publish(job,['web'],run_ai=True,automatic_models=MODELS)
    assert result['qa']['passed'] and result['quick']['retained_count']==1
    assert result['quick']['applied_count']==0
    assert '기준 값은 10이며 함수 호출이 가능합니다.' in Path(result['outputs']['web']).read_text(encoding='utf-8')


def test_grounded_technical_change_requires_final_support(tmp_path,monkeypatch):
    original='이 함수는 순서를 변경하지 않습니다.'
    evidence='설계 명세에 따르면 이 함수는 순서를 변경합니다.'
    job=job_for(tmp_path,[original,evidence])
    def response(role,p):
        if role=='final_review':
            return {'accept':True,'meaning_preserved':False,'evidence_supported':True,'reason':'제공 명세와 일치'}
        target=p['candidate']
        changed=role=='technical_review' and p['original']==original
        if changed:target='이 함수는 순서를 변경합니다.'
        return {'text':target,'technical_change':changed,'evidence':[{'quote':evidence}] if changed else [],'reason':'명세 대조'}
    fixtures(monkeypatch,response)
    _,report=automatic.run(job,MODELS)
    assert report['applied_count']==1
    assert report['applied'][0]['suggested_text']=='이 함수는 순서를 변경합니다.'


def test_question_region_is_not_sent_to_ai(tmp_path,monkeypatch):
    job=job_for(tmp_path,['연습문제','옳지 않은 설명을 고르시오.','보기의 내용은 오류입니다.'])
    calls=fixtures(monkeypatch)
    _,report=automatic.run(job,MODELS)
    assert not calls and not report['applied']


def test_malformed_reply_is_failure_not_completed(tmp_path,monkeypatch):
    job=job_for(tmp_path)
    fixtures(monkeypatch,lambda *_:{'suggestions':[]})
    with pytest.raises(RuntimeError,match='응답 형식'):
        publish(job,['web'],run_ai=True,automatic_models=MODELS)
    assert job.last_result() is None
    assert not list((job.work/'automatic-cache').glob('*.json'))


def test_model_change_invalidates_cache(tmp_path,monkeypatch):
    job=job_for(tmp_path)
    calls=fixtures(monkeypatch)
    automatic.run(job,MODELS)
    changed={**MODELS,'proofreading':('new-model','none')}
    automatic.run(job,changed)
    assert any(model=='new-model' for model,_ in calls)
    assert len(calls)==4  # Same candidate can reuse unchanged reviewers.


def test_change_manager_rollback_does_not_touch_manual_records(tmp_path,monkeypatch):
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication, QDialog, QTableWidget, QPushButton
    from app.gui.automatic_changes import show_changes
    job=job_for(tmp_path)
    fixtures(monkeypatch)
    _,report=automatic.run(job,MODELS)
    app=QApplication.instance() or QApplication([])
    checked=[]
    def interact():
        dialog=next(w for w in app.topLevelWidgets() if isinstance(w,QDialog) and w.windowTitle()=='자동 변경 관리')
        dialog.findChild(QTableWidget).selectRow(0)
        next(b for b in dialog.findChildren(QPushButton) if b.text()=='선택 문단 되돌리기').click()
        checked.append(True)
        dialog.accept()
    QTimer.singleShot(20,interact)
    show_changes(None,job,report,tmp_path)
    assert checked
    assert json.loads((job.work/'automatic-exclusions.json').read_text(encoding='utf-8'))==[report['applied'][0]['block_id']]
    assert job.workflow().items()==[]
