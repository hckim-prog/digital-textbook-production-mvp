import copy
from dataclasses import asdict
from pathlib import Path
from threading import Event
import json

from docx import Document
from PIL import Image
import pytest

from app.controllers.production import Production
from app.controllers.publishing import publish, preflight
from core.ai import service
from core.ai.workflow import write_json
from core.cancellation import OperationCancelled
from core.master import Block, Master
from core.manuscript import learning
from core.qa import editorial

MODELS = {role: (role, 'none') for role in ('proofreading','technical_review','final_review')}
TEXT = '함수는 이름을 통해 호출하며 입력을 받아 결과를 반환합니다.'


def job_for(tmp_path, placeholder=True):
    source = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph('Chapter 1 함수')
    doc.add_paragraph('1.1 함수 호출')
    doc.add_paragraph(TEXT)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'int main() {\n\treturn 0;\n}'
    picture = tmp_path / 'image.png'
    Image.new('RGB', (24, 18), 'blue').save(picture)
    doc.add_picture(str(picture))
    if placeholder:
        doc.add_paragraph('연습문제')
        doc.add_paragraph('(추후 제공)')
    doc.save(source)
    job = Production(tmp_path / 'project', source)
    job.analyze()
    return job


def fake_ai(monkeypatch, transform=None, cancelled=None):
    calls = []
    monkeypatch.setattr(service, 'model_config', lambda *_: {'enabled':True, 'reasoning_options':['none']})
    monkeypatch.setattr(service, 'record', lambda *a: None)
    def respond(root, model, reasoning, prompt, **options):
        payload = json.loads(prompt.split('자료:\n')[1])
        calls.append((prompt, payload))
        if 'questions' in payload:
            data = {'accepted':True, 'reason':'질문과 정답을 원고의 함수 설명과 대조함'}
        elif 'sources' in payload:
            source = next(r for r in payload['sources'] if r['text'] == TEXT)
            count = 2 if payload['type'] == 'section' else 5
            data = {'questions':[{'question': f'함수 호출의 의미를 설명하세요 ({i+1}).', 'answer':'함수 이름을 통해 호출합니다.',
                    'explanation':'원고는 함수가 입력을 받아 결과를 반환한다고 설명합니다.',
                    'evidence':[{'block_id':source['id'],'quote':TEXT}]} for i in range(count)]}
        elif 'units' in payload:
            data = {'issues':[]}
        elif 'technical_change' in payload:
            data = {'accept':True, 'meaning_preserved':True, 'evidence_supported':False, 'reason':'의미 보존'}
        else:
            data = {'text':payload['candidate'], 'technical_change':False, 'reason':'유지', 'evidence':[]}
        if transform:
            data = transform(payload, data)
        if cancelled and 'sources' in payload and 'questions' not in payload:
            cancelled.set()
        return data, {'model':model}
    monkeypatch.setattr(service, '_response', respond)
    return calls


def test_rules_ignore_code_and_normal_prose_but_inspect_prose_cells():
    master = Master('시험','test','hash',[
        Block('code','code-block',text='// TODO: implement\n**ptr;'),
        Block('ok','paragraph',text='초안 작성 방법과 C++는 정상적인 표기이다.'),
        Block('table','table',rows=[['TODO','(추후 제공)']],cell_kinds=[['code-block','paragraph']]),
        Block('missing','paragraph',text='[FIGURE missing-image]')])
    issues = editorial.inspect(master)
    assert {i['block_id'] for i in issues if i['severity']=='error'} == {'table-r0c1','missing'}


def test_placeholder_blocks_all_output_and_preserves_last_success(tmp_path):
    job = job_for(tmp_path)
    with pytest.raises(ValueError, match='완결성 검사 미통과'):
        publish(job, ['web','pdf','epub'])
    assert job.last_result() is None
    assert not list(job.default_output_base.glob('*/HTML/index.html'))
    assert json.loads((job.work/'editorial-qa.json').read_text(encoding='utf-8'))['status']=='FAIL'


def test_grounded_questions_all_formats_and_cache(tmp_path, monkeypatch):
    job = job_for(tmp_path)
    original = job.source.read_bytes()
    baseline = asdict(job.master())
    calls = fake_ai(monkeypatch)
    result = publish(job, ['web','pdf','epub'], run_ai=True, automatic_models=MODELS, learning=True, editorial_ai=True)
    assert result['qa']['passed'], result['qa']
    assert result['learning']['question_count']==7
    assert result['editorial']['ai_audited'] and result['editorial']['passed']
    assert result['editorial']['status']=='REVIEW_REQUIRED'  # Image meaning is not verified.
    text = Path(result['outputs']['web']).read_text(encoding='utf-8')
    assert '(추후 제공)' not in text and '절 확인 활동' in text and '장말 연습문제' in text and '예시 정답과 해설' in text
    assert text.index('절 확인 활동') < text.index('장말 연습문제')
    output = Master.load(Path(result['folder'])/'reports/approved-master.json')
    changed = {i['block_id'] for i in result['learning']['replacements']}
    assert [asdict(b) for b in output.blocks if not b.id.startswith('learning-') and b.id not in changed] == [asdict(b) for b in job.master().blocks if b.id not in changed]
    count = len(calls)
    publish(job, ['web'], run_ai=True, automatic_models=MODELS, learning=True, editorial_ai=True)
    assert len(calls)==count
    assert job.source.read_bytes()==original and asdict(job.master())==baseline and job.workflow().items()==[]
    audit_payloads = [p for _,p in calls if 'units' in p]
    assert audit_payloads and 'int main' not in json.dumps(audit_payloads)


@pytest.mark.parametrize('failure',['evidence','review','malformed','incomplete'])
def test_learning_failures_never_publish(tmp_path, monkeypatch, failure):
    job = job_for(tmp_path)
    def transform(p, d):
        if 'questions' in d and failure=='evidence':
            d['questions'][0]['evidence'][0]['quote']='원고에 없는 근거 문장입니다.'
        if 'accepted' in d and failure=='review':
            d['accepted']=False
        if 'questions' in d and failure=='malformed':
            return {'questions':'invalid'}
        if 'questions' in d and failure=='incomplete':
            d['questions'][0]['answer']='(추후 제공)'
        return d
    fake_ai(monkeypatch, transform)
    with pytest.raises((ValueError,RuntimeError)):
        publish(job,['web'],run_ai=True,automatic_models=MODELS,learning=True)
    assert job.last_result() is None


def test_manual_placeholder_decision_is_not_replaced(tmp_path,monkeypatch):
    job = job_for(tmp_path)
    workflow=job.workflow();workflow.start()
    block = job.master().blocks[-1]
    path=workflow.folder/'suggestions.json'
    write_json(path,[{'id':'manual','block_id':block.id,'source_text':block.text,'suggested_text':block.text,
                     'status':'hold','role':'proofreading'}])
    original=path.read_bytes()
    fake_ai(monkeypatch)
    with pytest.raises(ValueError,match='보호된 미완성'):
        publish(job,['web'],run_ai=True,automatic_models=MODELS,learning=True)
    assert path.read_bytes()==original and job.last_result() is None


def test_definite_residual_typo_blocks_publication(tmp_path, monkeypatch):
    job=job_for(tmp_path,False)
    def transform(p,d):
        if 'units' in p:
            unit=next(u for u in p['units'] if u['text']==TEXT)
            return {'issues':[{'block_id':unit['id'],'category':'typo','quote':'함수',
                              'message':'고정 시험의 오류 후보','definite':True}]}
        return d
    fake_ai(monkeypatch,transform)
    with pytest.raises(ValueError,match='완결성 검사 미통과'):
        publish(job,['web'],run_ai=True,automatic_models=MODELS,editorial_ai=True)
    assert job.last_result() is None


def test_unanchored_audit_response_is_rejected(tmp_path,monkeypatch):
    job=job_for(tmp_path,False)
    def transform(p,d):
        if 'units' in p:
            return {'issues':[{'block_id':'invented','category':'typo','quote':'없음','message':'오류','definite':True}]}
        return d
    fake_ai(monkeypatch,transform)
    with pytest.raises(RuntimeError,match='위치·근거'):
        publish(job,['web'],run_ai=True,automatic_models=MODELS,editorial_ai=True)


def test_learning_cancellation_reuses_finished_request(tmp_path,monkeypatch):
    job=job_for(tmp_path,False)
    stop=Event()
    calls=fake_ai(monkeypatch,cancelled=stop)
    nodes=preflight(job)['nodes']
    with pytest.raises(OperationCancelled):
        learning.generate(job,job.master(),nodes,MODELS,cancelled=stop.is_set)
    first=calls[0][0]
    stop.clear()
    calls2=fake_ai(monkeypatch)
    completed=learning.generate(job,job.master(),nodes,MODELS,cancelled=stop.is_set)
    assert completed['question_count']==7 and all(prompt != first for prompt,_ in calls2)


def test_addition_baseline_change_is_rejected(tmp_path,monkeypatch):
    job=job_for(tmp_path,False);fake_ai(monkeypatch)
    master=job.master()
    report=learning.generate(job,master,preflight(job)['nodes'],MODELS)
    modified=copy.deepcopy(master);modified.blocks[2].text+=' 변경'
    with pytest.raises(ValueError,match='기준 원고'):
        learning.apply(modified,report)


def test_without_ai_reports_unverified_not_pass(tmp_path):
    job=job_for(tmp_path,False)
    result=publish(job,['web'])
    assert result['editorial']['status']=='REVIEW_REQUIRED' and not result['editorial']['ai_audited']


def test_learning_options_require_explicit_ai_mode(tmp_path):
    job=job_for(tmp_path)
    with pytest.raises(ValueError,match='자동 수정'):
        publish(job,['web'],learning=True)


def test_repeated_chapter_label_is_skipped_without_changing_outline(tmp_path, monkeypatch):
    from core.manuscript.structure import propose_rules
    job = job_for(tmp_path, False)
    master = job.master()
    master.blocks.insert(0, Block('cover-label', 'paragraph', text='Chapter 1'))
    nodes = propose_rules(master)
    before = copy.deepcopy(nodes)
    calls = fake_ai(monkeypatch)
    result = learning.generate(job, master, nodes, MODELS)
    assert len(result['skipped_units']) == 1
    assert result['skipped_units'][0]['title'] == 'Chapter 1'
    assert result['question_count'] == 7 and nodes == before
    assert not any(p.get('id') == 'outline-0001' for _, p in calls)
    master.outline = nodes
    assert not any(i['block_id'] == 'cover-label' and i['code'] == 'missing_learning' for i in editorial.inspect(master))


@pytest.mark.parametrize('body', [
    Block('body', 'paragraph', text='짧은 설명.'),
    Block('body', 'code-block', text='return 0;'),
    Block('body', 'table', rows=[['값', '의미']]),
])
def test_short_prose_code_and_tables_are_learning_content(body):
    from core.manuscript.structure import node
    master = Master('시험', 's', 'hash', [Block('title','paragraph',text='Chapter 1'), body])
    assert len(learning.units(master, [node(master.blocks[0],1)])) == 1


def test_generated_heading_anchor_keeps_real_prose():
    from core.manuscript.structure import node
    master = Master('시험','s','hash',[Block('body','paragraph',text=TEXT)])
    assert len(learning.units(master,[node(master.blocks[0],1,'생성된 장 제목')])) == 1


def test_title_blank_and_placeholder_skip_requests_but_fail_completion(tmp_path,monkeypatch):
    from core.manuscript.structure import node
    job=job_for(tmp_path)
    master=Master('시험','s','hash',[Block('title','heading',text='Chapter 1'),
        Block('blank','paragraph',text='  '), Block('label','paragraph',text='연습문제'),
        Block('pending','paragraph',text='(추후 제공)')])
    calls=fake_ai(monkeypatch)
    result=learning.generate(job,master,[node(master.blocks[0],1)],MODELS)
    assert not calls and result['question_count']==0 and len(result['skipped_units'])==1
    assert not editorial.report(master)['passed']


def test_empty_response_still_fails_with_unit_and_count(tmp_path,monkeypatch):
    job=job_for(tmp_path,False)
    fake_ai(monkeypatch,lambda p,d: {'questions':[]} if 'sources' in p and 'questions' not in p else d)
    with pytest.raises(RuntimeError,match='Chapter 1 함수.*반환된 문제: 0개'):
        learning.generate(job,job.master(),preflight(job)['nodes'],MODELS)
