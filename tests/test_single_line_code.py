import pytest
from docx import Document
from core.manuscript.docx_reader import read_docx
from core.manuscript.code_blocks import scan_code_sources
from core.export.outputs import web, pdf, epub_file
from core.qa.code_fidelity import check_code_fidelity, check_output_codes


@pytest.mark.parametrize('code', ['cout << getLarger(10.0, 20.5) << endl;', '\tcout << 20;  ', '    return 0;'])
def test_single_line_table_code_exact_preservation(tmp_path, code):
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = code
    source = tmp_path / 'one.docx'
    doc.save(source)
    region = scan_code_sources(Document(source))[0]
    assert region.block_id == 'b00001-r0c0'  # Existing saved master IDs stay usable.
    assert region.text == code and region.table_cell == (0, 0)
    master = read_docx(source, tmp_path / 'work')
    assert master.blocks[0].kind == 'table'
    assets = tmp_path / 'work/assets'
    outputs = {'web': web(master, assets, tmp_path / 'html'),
               'pdf': pdf(master, assets, tmp_path / 'pdf'),
               'epub': epub_file(master, assets, tmp_path / 'epub')}
    assert check_code_fidelity(source, master, outputs['web'])['passed']
    assert check_output_codes(source, master, outputs)['passed']
    master.blocks[0].rows[0][0] += ' '
    assert not check_code_fidelity(source, master, outputs['web'])['passed']


def test_code_keyword_in_prose_not_promoted(tmp_path):
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'cout 또는 std::vector를 설명하는 문장입니다.'
    source = tmp_path / 'prose.docx'
    doc.save(source)
    master = read_docx(source, tmp_path / 'work')
    assert not scan_code_sources(Document(source))
    assert master.blocks[0].cell_kinds == [['paragraph']]


def test_tall_table_row_can_span_pdf_pages(tmp_path):
    from app.controllers.production import Production
    doc = Document()
    doc.add_heading('긴 표 시험', 1)
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    for i in range(90):
        cell.add_paragraph(f'{i:03d}번 설명 문장입니다. 긴 표도 원문 순서를 보존해야 합니다.')
    source = tmp_path / 'tall.docx'
    doc.save(source)
    job = Production(tmp_path / 'project', source)
    job.analyze()
    result = job.build(['web', 'pdf', 'epub'])
    assert result['qa']['passed'], result['qa']['checks']
