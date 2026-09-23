import copy
import json
import pytest
from docx import Document
from core.ai.quick import allowed, quick_master
from core.master import Master, Block
from app.controllers.production import Production
from core.ai.workflow import write_json


def proposal(status='pending', target='할 수 있다.'):
    return dict(id='q1', block_id='b1', source_text='할수 있다.', suggested_text=target,
                level='A', role='proofreading', status=status, reason='fixture')


def test_allowlist_does_not_allow_paraphrase_numbers_or_code_whitespace():
    assert allowed('할수 있다.', '할 수 있다.')
    assert not allowed('값은 5이다.', '값은 6이다.')
    assert not allowed('할수 있다.', '가능하다.')
    assert not allowed('a  b', 'a b')
    assert not allowed('몇일변수', '며칠변수')


@pytest.mark.parametrize('status', ['approved', 'rejected', 'hold', 'outdated'])
def test_existing_decisions_untouched(status):
    m=Master('t','s','hash',[Block('b1','paragraph','할수 있다.')])
    items=[proposal(status)]
    before=copy.deepcopy(items)
    out, report=quick_master(m,items)
    assert items==before and out.blocks[0].text==m.blocks[0].text
    assert report['applied_count']==0


@pytest.mark.parametrize('kind', ['code-block','table','figure'])
def test_protected_blocks(kind):
    m=Master('t','s','hash',[Block('b1',kind,'할수 있다.')])
    out,report=quick_master(m,[proposal()])
    assert out==m and report['applied_count']==0


def test_conflicts_and_technical_proposals_kept():
    m=Master('t','s','hash',[Block('b1','paragraph','할수 있다.')])
    a=proposal();b={**a,'id':'q2'}
    assert quick_master(m,[a,b])[1]['applied_count']==0
    b['role']='technical_review'
    assert quick_master(m,[b])[1]['applied_count']==0


def test_export_only_changes_and_undo(tmp_path):
    source=tmp_path/'sample.docx';doc=Document();doc.add_paragraph('할수 있다.');doc.save(source)
    original=source.read_bytes()
    job=Production(tmp_path/'project',source);master=job.analyze();workflow=job.workflow();workflow.start()
    item=proposal();item['block_id']=master.blocks[0].id
    write_json(workflow.folder/'suggestions.json',[item])
    decisions=(workflow.folder/'suggestions.json').read_bytes()
    first=job.build(['web'],quick=True)
    assert first['qa']['passed'] and first['quick']['applied_count']==1
    after=Master.load(__import__('pathlib').Path(first['folder'])/'reports/approved-master.json')
    assert after.blocks[0].text=='할 수 있다.'
    second=job.build(['web'],quick=False)
    after=Master.load(__import__('pathlib').Path(second['folder'])/'reports/approved-master.json')
    assert after.blocks[0].text=='할수 있다.'
    assert source.read_bytes()==original
    assert (workflow.folder/'suggestions.json').read_bytes()==decisions


def test_quick_gui_runs_only_proofreading_and_builds(tmp_path, monkeypatch):
    import shutil
    import time
    from PySide6.QtWidgets import QApplication
    from PySide6.QtCore import QSettings
    import app.gui.window as gui
    from core.ai import workflow as engine
    app=QApplication.instance() or QApplication([])
    root=tmp_path/'project';shutil.copytree(__import__('pathlib').Path(__file__).resolve().parents[1]/'config',root/'config')
    source=tmp_path/'sample.docx';doc=Document();doc.add_paragraph('할수 있다.');doc.save(source)
    prefs=QSettings(str(tmp_path/'prefs.ini'),QSettings.IniFormat)
    monkeypatch.setattr(gui,'QSettings',lambda *_:prefs)
    roles=[]
    def fake(root,master,model,reasoning,limit,progress,role,cancelled=None):
        roles.append(role)
        item=proposal();item['block_id']=master.blocks[0].id;item['model']=model
        return [item],[dict(success=True,block_id=master.blocks[0].id,input_tokens=1,output_tokens=1)]
    monkeypatch.setattr(engine,'proofread',fake)
    win=gui.MainWindow(root)
    win.depth_existing.setChecked(True)
    win.source_edit.setText(str(source));win.prepare_timer.stop()
    job=Production(root,source);master=job.analyze()
    from core.manuscript.structure import node
    job.structure().save([node(master.blocks[0],1,'시험 장')],approved=True)
    win._analysis_done(job,master)
    for key,box in win.format_boxes.items():box.setChecked(key=='web')
    assert win.analysis_panel.isHidden() and not win.detail_toggle.isChecked()
    win.build()
    deadline=time.monotonic()+15
    while (win.busy or any(t.isRunning() for t in win.threads)) and time.monotonic()<deadline:
        app.processEvents();time.sleep(.01)
    assert not win.busy and roles==['proofreading']
    assert win.last_output_folder and win.changes_button.isEnabled()
    report=json.loads((win.last_output_folder/'reports/quick-changes.json').read_text(encoding='utf-8'))
    assert report['applied_count']==1
    assert job.workflow().items()[0]['status']=='pending'
    win.close()
