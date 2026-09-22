import copy
from dataclasses import asdict
import json

import pytest
from docx import Document
from lxml import html
from PySide6.QtWidgets import QApplication, QMessageBox

from core.master import Block, Master, file_hash
from core.manuscript.structure import StructureStore, ai_payload, node, propose_ai, propose_rules, validate
from core.manuscript.docx_reader import read_docx
from app.controllers.production import Production


def sample():
    return Master("시험 교재", "source.docx", "hash", [
        Block("b1", "paragraph", "들어가는 말"),
        Block("b2", "paragraph", "1.1. 첫 번째 절"),
        Block("b3", "paragraph", "1.1.1. 첫 소단원"),
        Block("b4", "code-block", "#include <iostream>\n\nint main() {\n\treturn 0;\n}"),
        Block("b5", "paragraph", "1.2. 두 번째 절"),
        Block("b6", "heading", "💡 팁: 코드 입력", "Heading 3"),
        Block("b7", "paragraph", "끝")])


def test_ranges_and_exact_preservation(tmp_path):
    master = sample()
    before = copy.deepcopy(asdict(master))
    nodes = propose_rules(master)
    assert [n["level"] for n in nodes] == [1, 2, 3, 2]
    assert nodes[1]["end_block_id"] == "b4"
    assert nodes[2]["parent_id"] == nodes[1]["id"]
    store = StructureStore(tmp_path, master)
    store.save(nodes)
    with pytest.raises(ValueError, match="승인"):
        store.apply(master)
    store.save(nodes, True)
    approved = store.apply(master)
    assert asdict(master) == before
    assert asdict(approved)["blocks"] == before["blocks"]
    approved.save(tmp_path / "master.json")
    assert asdict(Master.load(tmp_path / "master.json")) == asdict(approved)
    master.blocks[0].text = "변경"
    with pytest.raises(ValueError, match="변경"):
        store.load()


@pytest.mark.parametrize("change", ["skip", "missing", "duplicate", "reorder", "code", "prefix"])
def test_invalid_structures_rejected(change):
    master = sample()
    nodes = propose_rules(master)
    if change == "skip":
        nodes[1]["level"] = 3
    elif change == "missing":
        nodes[1]["start_block_id"] = "unknown"
    elif change == "duplicate":
        nodes.insert(2, copy.deepcopy(nodes[1]))
    elif change == "reorder":
        nodes[1], nodes[3] = nodes[3], nodes[1]
    elif change == "code":
        nodes[2]["start_block_id"] = "b4"
    else:
        nodes[0]["start_block_id"] = "b2"
    with pytest.raises(ValueError):
        validate(master, nodes)


def test_ai_protection_and_response_validation(tmp_path, monkeypatch):
    import core.manuscript.structure as module
    master = sample()
    payload = ai_payload(master)
    assert "#include" not in payload and "return 0" not in payload
    calls = []
    monkeypatch.setattr(module, "record", lambda *args: calls.append(args))
    monkeypatch.setattr(module, "_response", lambda *args: ({"nodes": propose_rules(master)}, {"model": "fake"}))
    assert len(propose_ai(tmp_path, master, "fake", "none")) == 4
    monkeypatch.setattr(module, "_response", lambda *args: ({"nodes": [{"level": 1, "start_block_id": "invented"}]}, {"model": "fake"}))
    with pytest.raises(ValueError):
        propose_ai(tmp_path, master, "fake", "none")
    assert calls[-1][-2] is False


def test_docx_to_all_formats_and_gui(tmp_path, monkeypatch):
    source = tmp_path / "교재.docx"
    doc = Document()
    doc.add_paragraph("시작 설명")
    doc.add_heading("첫 절", level=2)
    doc.add_heading("첫 소단원", level=3)
    code = '#include <iostream>\n\nint main() {\n\treturn 0;\n}'
    doc.add_table(rows=1, cols=1).cell(0, 0).text = code
    doc.add_heading("다음 절", level=2)
    doc.add_paragraph("마지막 문장")
    doc.save(source)
    digest = file_hash(source)
    job = Production(tmp_path / "project", source)
    master = job.analyze()
    nodes = propose_rules(master)
    assert [n['level'] for n in nodes] == [1, 2, 3, 2]
    assert master.blocks[1].outline_level == 2
    app = QApplication.instance() or QApplication([])
    from app.gui.structure_dialog import StructureDialog
    dialog = StructureDialog(job, "fake", "none")
    dialog.rules()
    assert dialog.tree.topLevelItem(0).childCount() == 2
    dialog.tree.setCurrentItem(dialog.tree.topLevelItem(0).child(0).child(0))
    assert code in dialog.preview.toPlainText()
    dialog.approve()
    result = job.build(["web", "pdf", "epub"])
    assert result["qa"]["passed"], result["qa"]["checks"]
    tree = html.parse(result["outputs"]["web"])
    assert len(tree.xpath("//h1")) == 1
    assert len(tree.xpath("//h2")) == 2
    assert len(tree.xpath("//h3")) == 1
    assert tree.xpath("//pre/code")[0].text == code
    assert file_hash(source) == digest
    approved = Master.load(job.work / "approved-master.json")
    assert [asdict(b) for b in approved.blocks] == [asdict(b) for b in master.blocks]
