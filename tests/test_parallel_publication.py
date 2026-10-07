import copy
import json
from pathlib import Path
from threading import Barrier, Event, Lock, current_thread, enumerate as threads
from types import SimpleNamespace

import pytest

from core.ai import automatic, service
from core.ai.workflow import write_json
from core.cancellation import OperationCancelled
from core.master import Block, Master
from core.qa import editorial
from test_automatic import job_for, MODELS, fixtures


class APITimeoutError(RuntimeError):
    """Local transport fixture; no network or SDK request is made."""


def reply(p, changed=False):
    return {'text': p['candidate'].replace('실제 호출', '실제로 호출') if changed else p['candidate'],
            'technical_change': False, 'evidence': [], 'reason': '고정 응답'}


def test_parallel_groups_reduce_requests_and_finalize_in_source_order(tmp_path, monkeypatch):
    job = job_for(tmp_path, [f'실제 호출할 함수의 동작을 설명합니다 {i}.' for i in range(6)])
    fixtures(monkeypatch)
    barrier, lock = Barrier(2), Lock()
    calls, finals = [], []
    active, peak = 0, 0

    def respond(root, model, effort, prompt, **options):
        nonlocal active, peak
        p = json.loads(prompt.split('자료:\n')[1])
        with lock:
            active += 1
            peak = max(peak, active)
            calls.append((model, p))
        try:
            if model == 'proofreading':
                barrier.wait(timeout=5)  # Fails if the two groups actually run serially.
            if model == 'final_review':
                finals.append((p['block_id'], current_thread().name))
                return {'accept': True, 'meaning_preserved': True, 'evidence_supported': False}, {}
            assert options['timeout'] == 120
            # Reversed responses must still map by ID, never by returned position.
            return {'results': [{'block_id': item['block_id'], 'response': reply(item, True)}
                                for item in reversed(p['items'])]}, {}
        finally:
            with lock:
                active -= 1

    monkeypatch.setattr(service, '_response', respond)
    baseline, eligible, _ = automatic.plan(job)
    result, report = automatic.run(job, MODELS)
    assert peak == 2
    assert len(calls) == 10  # Four batch requests + six ordered final checks, instead of 18.
    assert [id_ for id_, _ in finals] == [b.id for b in eligible]
    assert all(name == current_thread().name for _, name in finals)
    assert [i['block_id'] for i in report['applied']] == [b.id for b in eligible]
    automatic.verify_result(baseline, result, report)
    assert all('int main' not in json.dumps(p) for _, p in calls)
    monkeypatch.setattr(service, '_response', lambda *a, **k: pytest.fail('cached response requested again'))
    again, saved = automatic.run(job, MODELS)
    assert again == result and saved == report


@pytest.mark.parametrize('invalid', ['duplicate', 'missing', 'foreign', 'malformed'])
def test_batch_envelope_errors_save_no_partial_responses(tmp_path, monkeypatch, invalid):
    job = job_for(tmp_path, ['일반 문장 첫째입니다.', '일반 문장 둘째입니다.', '일반 문장 셋째입니다.'])
    fixtures(monkeypatch)
    def respond(root, model, effort, prompt, **options):
        items = json.loads(prompt.split('자료:\n')[1])['items']
        rows = [{'block_id': p['block_id'], 'response': reply(p)} for p in items]
        if invalid == 'duplicate': rows[-1] = copy.deepcopy(rows[0])
        if invalid == 'missing': rows.pop()
        if invalid == 'foreign': rows[-1]['block_id'] = 'invented'
        if invalid == 'malformed': rows[-1]['response'] = {'text': '불완전 응답'}
        return {'results': rows}, {}
    monkeypatch.setattr(service, '_response', respond)
    with pytest.raises(RuntimeError):
        automatic.run(job, MODELS)
    assert not list((job.work / 'automatic-cache').glob('*.json'))
    assert not (job.work / 'automatic-last.json').exists()
    assert not any(t.name.startswith('prose-review') for t in threads())


def test_cancel_drains_two_workers_saves_replies_and_resumes(tmp_path, monkeypatch):
    job = job_for(tmp_path, [f'일반 문장 설명입니다 {i}.' for i in range(6)])
    fixtures(monkeypatch)
    barrier, stop, calls = Barrier(2), Event(), []
    def respond(root, model, effort, prompt, **options):
        p = json.loads(prompt.split('자료:\n')[1])
        calls.append(model)
        barrier.wait(timeout=5)
        stop.set()
        return {'results': [{'block_id': item['block_id'], 'response': reply(item)} for item in p['items']]}, {}
    monkeypatch.setattr(service, '_response', respond)
    with pytest.raises(OperationCancelled):
        automatic.run(job, MODELS, cancelled=stop.is_set)
    assert calls == ['proofreading', 'proofreading']
    assert len(list((job.work / 'automatic-cache').glob('*.json'))) == 6
    assert not (job.work / 'automatic-last.json').exists()
    assert not any(t.name.startswith('prose-review') for t in threads())
    calls = fixtures(monkeypatch)
    _, report = automatic.run(job, MODELS)
    assert [role for role, _ in calls] == ['technical_review', 'technical_review']
    assert report['unchanged_count'] == 6


def test_worker_failure_is_failure_not_user_cancellation(tmp_path, monkeypatch):
    job = job_for(tmp_path, [f'일반 문장 설명입니다 {i}.' for i in range(6)])
    fixtures(monkeypatch)
    barrier, failure = Barrier(2), Event()
    first_id = automatic.plan(job)[1][0].id
    def respond(root, model, effort, prompt, **options):
        p = json.loads(prompt.split('자료:\n')[1])
        barrier.wait(timeout=5)
        if p['items'][0]['block_id'] != first_id:
            failure.set()
            raise ValueError('시험 요청 실패')
        assert failure.wait(5)
        return {'results': [{'block_id': item['block_id'], 'response': reply(item)} for item in p['items']]}, {}
    monkeypatch.setattr(service, '_response', respond)
    with pytest.raises(RuntimeError, match='묶음 교정 요청 실패') as error:
        automatic.run(job, MODELS)
    assert not isinstance(error.value, OperationCancelled)
    assert not (job.work / 'automatic-last.json').exists()


def test_legacy_single_cache_keys_and_partial_hits_are_reused(tmp_path, monkeypatch):
    job = job_for(tmp_path, ['일반 문장 첫째입니다.', '일반 문장 둘째입니다.', '일반 문장 셋째입니다.'])
    calls = fixtures(monkeypatch)
    states = list(automatic._groups(automatic.plan(job)[1]))[0]
    # Populate the exact single-request path retained from previous releases.
    s = states[0]
    payload = {'block_id': s['block'].id, 'original': s['block'].text,
               'candidate': s['candidate'], 'context': s['context']}
    for role in automatic.ROLES[:2]:
        automatic._request(job, role, *MODELS[role], payload, None)
    before = {p: p.read_bytes() for p in (job.work / 'automatic-cache').glob('*.json')}
    calls.clear()
    _, report = automatic.run(job, MODELS)
    assert len(calls) == 2 and report['unchanged_count'] == 3
    assert all(len(p['items']) == 2 for _, p in calls)
    assert all(payload['block_id'] not in [i['block_id'] for i in p['items']] for _, p in calls)
    assert all(p.read_bytes() == contents for p, contents in before.items())


def test_editorial_timeout_labels_small_chunks_and_resume(tmp_path, monkeypatch):
    job = SimpleNamespace(work=tmp_path/'work', root=tmp_path, source=Path('fixture.docx'))
    fixtures(monkeypatch)
    raw = '가' * 3995 + '경계에서 보존할 문장입니다.' + '나' * 5000
    master = Master('test', 'source', 'hash', [Block('prose', 'paragraph', text=raw),
        Block('code', 'code-block', text='\tcode\n\n  whitespace')])
    chunks = list(editorial.audit_units(master))
    assert chunks[0]['text'] + ''.join(c['text'][200:] for c in chunks[1:]) == raw
    assert any('경계에서 보존할 문장입니다.' in c['text'] for c in chunks)
    calls = []
    def respond(root, model, effort, prompt, **options):
        p = json.loads(prompt.split('자료:\n')[1])
        assert options == {'timeout': 120}
        assert sum(len(u['text']) for u in p['units']) <= 4000
        assert all(u['id'] == 'prose' for u in p['units'])
        calls.append(p['id'])
        if len(calls) == 2:
            raise APITimeoutError('local timeout fixture')
        return {'issues': []}, {}
    monkeypatch.setattr(service, '_response', respond)
    with pytest.raises(RuntimeError, match='최종 문장 검사 2/3.*대기시간'):
        editorial.audit(job, master, ('final_review','none'))
    assert len(list((job.work/'publication-cache').glob('*.json'))) == 1
    assert editorial.audit(job, master, ('final_review','none')) == []
    assert calls == ['batch-0', 'batch-1', 'batch-1', 'batch-2']
    assert editorial.audit(job, master, ('final_review','none')) == []
    assert len(calls) == 4


def test_gui_failure_keeps_progress_and_offers_resume(tmp_path, monkeypatch):
    import shutil
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication
    import app.gui.window as gui
    app = QApplication.instance() or QApplication([])
    root = tmp_path/'project'
    shutil.copytree(Path(__file__).resolve().parents[1]/'config', root/'config')
    monkeypatch.setattr(gui, 'QSettings', lambda *_: QSettings(str(tmp_path/'prefs.ini'), QSettings.IniFormat))
    monkeypatch.setattr(gui.QMessageBox, 'warning', lambda *a: None)
    win = gui.MainWindow(root)
    try:
        win.is_building = win.busy = True
        win.progress.setValue(73)
        win._failed('최종 문장 검사 2/3 실패')
        assert win.progress.value() == 73 and win.resume_build and not win.busy
        assert win.build_button.text() == '이어서 제작'
    finally:
        win.close()


def test_corrected_word_symbol_survives_all_outputs(tmp_path, monkeypatch):
    from docx import Document
    from app.controllers.production import Production
    from app.controllers.publishing import publish
    source = tmp_path/'symbols.docx'
    doc = Document()
    doc.add_paragraph('Chapter 1 시작')
    doc.add_paragraph('1.1 본문')
    p = doc.add_paragraph('실제 호출할 식은 x ')
    p.add_run('\uf02d').font.name = 'Symbol'
    p.add_run(' y입니다.')
    doc.save(source)
    original = source.read_bytes()
    job = Production(tmp_path/'project', source)
    job.analyze()
    fixtures(monkeypatch)
    result = publish(job, ['web','pdf','epub'], run_ai=True, automatic_models=MODELS)
    assert result['qa']['passed'] and result['quick']['applied_count'] == 1
    rendered = Master.load(Path(result['folder'])/'reports/approved-master.json')
    block = next(b for b in rendered.blocks if '실제로 호출' in b.text)
    assert ''.join(i.get('text','') for i in block.inlines) == block.text
    assert any(i.get('font') == 'Symbol' and i['text'] == '\uf02d' for i in block.inlines)
    assert source.read_bytes() == original
    saved = (job.work/'automatic-last.json').read_bytes()
    monkeypatch.setattr(service, '_response', lambda *a, **k: pytest.fail('cache should be reused'))
    again = publish(job, ['web'], run_ai=True, automatic_models=MODELS)
    assert again['qa']['passed'] and (job.work/'automatic-last.json').read_bytes() == saved
